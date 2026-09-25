import csv
import io
import json
import logging
from datetime import UTC, datetime

logger = logging.getLogger("2to-eos.export")


def export_to_csv(
    data: list[dict],
    columns: list[str] | None = None,
    filename: str = "export.csv",
) -> tuple[bytes, str]:
    if not data:
        return b"", "text/csv"
    
    if not columns:
        columns = list(data[0].keys())
    
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=columns, extrasaction="ignore")
    writer.writeheader()
    
    for row in data:
        writer.writerow({k: str(v) if v is not None else "" for k, v in row.items()})
    
    content = output.getvalue().encode("utf-8-sig")
    return content, "text/csv"


def export_to_json(
    data: list[dict] | dict,
    pretty: bool = True,
) -> tuple[bytes, str]:
    if isinstance(data, list):
        content = json.dumps(data, indent=2 if pretty else None, ensure_ascii=False, default=str)
    else:
        content = json.dumps(data, indent=2 if pretty else None, ensure_ascii=False, default=str)
    
    return content.encode("utf-8"), "application/json"


def export_to_xlsx(
    data: list[dict],
    columns: list[str] | None = None,
    sheet_name: str = "Sheet1",
    headers: dict[str, str] | None = None,
) -> tuple[bytes, str]:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError:
        logger.warning("openpyxl not installed, falling back to CSV")
        return export_to_csv(data, columns)
    
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    
    if not data:
        return io.BytesIO(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    
    if not columns:
        columns = list(data[0].keys())
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2B5797", end_color="2B5797", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = headers.get(col_name, col_name) if headers else col_name
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    for row_idx, row_data in enumerate(data, 2):
        for col_idx, col_name in enumerate(columns, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            value = row_data.get(col_name)
            cell.value = value if value is not None else ""
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
    
    for col_idx, col_name in enumerate(columns, 1):
        max_length = len(headers.get(col_name, col_name) if headers else col_name)
        for row_idx in range(2, len(data) + 2):
            cell_value = str(ws.cell(row=row_idx, column=col_idx).value or "")
            max_length = max(max_length, len(cell_value))
        ws.column_dimensions[chr(64 + col_idx) if col_idx <= 26 else "A"].width = min(max_length + 2, 50)
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def export_to_pdf(
    data: list[dict],
    title: str = "Report",
    columns: list[str] | None = None,
    headers: dict[str, str] | None = None,
    orientation: str = "landscape",
) -> tuple[bytes, str]:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.pagesizes import landscape as rl_landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        
        page_size = rl_landscape(A4) if orientation == "landscape" else A4
        
        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=page_size)
        
        styles = getSampleStyleSheet()
        elements = []
        
        title_style = ParagraphStyle(
            "CustomTitle",
            parent=styles["Heading1"],
            fontSize=16,
            spaceAfter=20,
            alignment=1,
        )
        elements.append(Paragraph(title, title_style))
        
        date_style = ParagraphStyle(
            "DateStyle",
            parent=styles["Normal"],
            fontSize=10,
            spaceAfter=10,
            alignment=1,
            textColor=colors.grey,
        )
        elements.append(Paragraph(f"Generated: {datetime.now(UTC).strftime('%Y-%m-%d %H:%M')}", date_style))
        elements.append(Spacer(1, 20))
        
        if not data:
            elements.append(Paragraph("No data available", styles["Normal"]))
            doc.build(elements)
            return output.getvalue(), "application/pdf"
        
        if not columns:
            columns = list(data[0].keys())
        
        display_headers = [headers.get(col, col) if headers else col for col in columns]
        table_data = [display_headers]
        
        for row in data:
            table_row = [str(row.get(col, "")) if row.get(col) is not None else "" for col in columns]
            table_data.append(table_row)
        
        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2B5797")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 10),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
            ("BACKGROUND", (0, 1), (-1, -1), colors.white),
            ("GRID", (0, 0), (-1, -1), 1, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F4F8")]),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 1), (-1, -1), 9),
            ("TOPPADDING", (0, 1), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 1), (-1, -1), 6),
        ]))
        
        elements.append(table)
        doc.build(elements)
        
        return output.getvalue(), "application/pdf"
    
    except ImportError:
        logger.warning("reportlab not installed, falling back to JSON")
        return export_to_json(data)


def format_export_response(
    data: list[dict],
    format: str,
    title: str = "Report",
    columns: list[str] | None = None,
    headers: dict[str, str] | None = None,
) -> tuple[bytes, str, str]:
    format_map = {
        "csv": ("text/csv", ".csv"),
        "json": ("application/json", ".json"),
        "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ".xlsx"),
        "pdf": ("application/pdf", ".pdf"),
    }
    
    if format == "csv":
        content, mime = export_to_csv(data, columns)
    elif format == "json":
        content, mime = export_to_json(data)
    elif format == "xlsx":
        content, mime = export_to_xlsx(data, columns, headers=headers)
    elif format == "pdf":
        content, mime = export_to_pdf(data, title, columns, headers)
    else:
        content, mime = export_to_csv(data, columns)
    
    ext = format_map.get(format, (".csv",))[1]
    timestamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
    filename = f"{title.lower().replace(' ', '_')}_{timestamp}{ext}"
    
    return content, mime, filename
